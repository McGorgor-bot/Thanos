import io, json, re, zipfile
from pathlib import Path
import pandas as pd
import streamlit as st
from .config import APP_VERSION, RESULT_COLUMNS
from .cad import collect_legend_graphics, _legend_display_name_map

def _add_image_fit_to_cell(ws, cell_ref, img_bytes, max_width_px=105, max_height_px=62, row_height_pt=68):
    """Attach the PNG to the target cell and fit it inside the cell bounds.

    openpyxl stores this as an OOXML one-cell anchor, so the image is tied to the
    specified cell rather than floating across multiple cells.
    """
    from openpyxl.drawing.image import Image as XLImage
    from PIL import Image as PILImage
    import tempfile, os
    if not img_bytes:
        return False, None
    # Use a real temporary PNG path. This is more reliable than transient BytesIO
    # streams when a workbook contains many images.
    tmp=tempfile.NamedTemporaryFile(suffix='.png', delete=False)
    tmp.write(img_bytes); tmp.flush(); tmp.close()
    try:
        with PILImage.open(tmp.name) as im:
            iw, ih=im.size
        if iw <= 0 or ih <= 0:
            return False, tmp.name
        scale=min(max_width_px/iw, max_height_px/ih, 1.0)
        w=max(1, int(round(iw*scale)))
        h=max(1, int(round(ih*scale)))
        img=XLImage(tmp.name)
        img.width=w; img.height=h
        ws.add_image(img, cell_ref)
        row=int(re.search(r'\d+',cell_ref).group())
        ws.row_dimensions[row].height=row_height_pt
        return True, tmp.name
    except Exception:
        return False, tmp.name


def enrich_legend_names_from_files(df, uploaded_files=None, exclude_legend=True):
    """Fill 'Nazwa z legendy' directly from each source DXF before XLSX export.
    This is deliberately done at export time so stale/cached result tables cannot
    leave the XLSX column empty.
    """
    out=df.copy() if isinstance(df,pd.DataFrame) else pd.DataFrame()
    if out.empty or not uploaded_files:
        return out
    if 'Nazwa z legendy' not in out.columns:
        out['Nazwa z legendy']=''
    file_maps={}
    for f in uploaded_files:
        try:
            if Path(f.name).suffix.lower() != '.dxf':
                continue
            raw=f.getvalue()
            file_maps[f.name]=_legend_display_name_map(raw, exclude_legend=exclude_legend)
        except Exception:
            file_maps[f.name]={}
    if not file_maps:
        return out
    for idx,row in out.iterrows():
        if str(row.get('Nazwa z legendy','') or '').strip():
            continue
        fname=str(row.get('Plik','') or '')
        mp=file_maps.get(fname)
        if not mp:
            # tolerate names with path prefixes/suffixes
            for k,v in file_maps.items():
                if Path(k).name==Path(fname).name or k==fname:
                    mp=v; break
        if not mp:
            continue
        keys=[row.get('Name',''),row.get('Misc',''),row.get('Nazwa bloku',''),row.get('Nazwa elementu','')]
        for key in keys:
            key=re.sub(r'\s+',' ',str(key or '')).strip().casefold()
            if key and key in mp:
                out.at[idx,'Nazwa z legendy']=mp[key]
                break
    return out


def apply_legend_names_to_results(df, legend_maps=None):
    """Apply cached legend maps to the visible result table."""
    out=df.copy() if isinstance(df,pd.DataFrame) else pd.DataFrame()
    if out.empty:
        return out
    if 'Nazwa z legendy' not in out.columns:
        out['Nazwa z legendy']=''
    legend_maps=legend_maps or {}
    for idx,row in out.iterrows():
        if str(row.get('Nazwa z legendy','') or '').strip():
            continue
        fname=str(row.get('Plik','') or '')
        mp=legend_maps.get(fname)
        if mp is None:
            for k,v in legend_maps.items():
                if Path(k).name==Path(fname).name:
                    mp=v; break
        if not mp: continue
        for key in (row.get('Name',''),row.get('Misc',''),row.get('Nazwa bloku',''),row.get('Nazwa elementu','')):
            key=re.sub(r'\s+',' ',str(key or '')).strip().casefold()
            if key in mp:
                out.at[idx,'Nazwa z legendy']=mp[key]
                break
    return out


def _convert_graphics_to_incell_xlsx(xlsx_bytes, cell_images):
    if not cell_images:
        return xlsx_bytes
    import zipfile, io, re
    zin=zipfile.ZipFile(io.BytesIO(xlsx_bytes),'r')
    files={n:zin.read(n) for n in zin.namelist()}
    for n in list(files):
        if n.startswith('xl/drawings/') or n.startswith('xl/media/'):
            del files[n]
    # RichData images belong to the Podsumowanie worksheet, not necessarily sheet1.
    # Resolve the worksheet path from workbook.xml + workbook relationships.
    sheet_path = None
    try:
        wbxml = files.get('xl/workbook.xml', b'').decode('utf-8')
        wbrels = files.get('xl/_rels/workbook.xml.rels', b'').decode('utf-8')
        m = re.search(r'<sheet\b[^>]*name=["\']Podsumowanie["\'][^>]*r:id=["\']([^"\']+)["\'][^>]*/?>', wbxml)
        if m:
            rid = m.group(1)
            rm = re.search(r'<Relationship\b[^>]*Id=["\']'+re.escape(rid)+r'["\'][^>]*Target=["\']([^"\']+)["\']', wbrels)
            if rm:
                target = rm.group(1).lstrip('/')
                sheet_path = target if target.startswith('xl/') else 'xl/' + target
    except Exception:
        sheet_path = None
    if sheet_path not in files:
        # Fallback: in older workbooks Podsumowanie was sheet2.
        if 'xl/worksheets/sheet2.xml' in files:
            sheet_path = 'xl/worksheets/sheet2.xml'
        elif 'xl/worksheets/sheet1.xml' in files:
            sheet_path = 'xl/worksheets/sheet1.xml'
        else:
            return xlsx_bytes
    sheet=files[sheet_path].decode('utf-8')
    for i,(cell,_) in enumerate(cell_images):
        # Keep the cell style used by the graphic column. Excel's Picture-in-Cell
        # rich value is represented internally by an error placeholder, but the
        # original style must remain attached to the cell.
        pat=rf'<c r="{re.escape(cell)}"([^>]*)>(?:.*?</c>)|<c r="{re.escape(cell)}"([^>]*)/>'
        m=re.search(pat,sheet,flags=re.S)
        style=''
        if m:
            attrs=(m.group(1) or m.group(2) or '')
            sm=re.search(r'\bs="(\d+)"',attrs)
            if sm:
                style=f' s="{sm.group(1)}"'
        repl=f'<c r="{cell}"{style} t="e" vm="{i+1}"><v>#VALUE!</v></c>'
        if m:
            sheet=sheet[:m.start()]+repl+sheet[m.end():]
        else:
            row=re.search(r'\d+',cell).group()
            sheet=re.sub(rf'(<row r="{row}"[^>]*>)',rf'\1{repl}',sheet,count=1)
    files[sheet_path]=sheet.encode()
    for i,(_,img) in enumerate(cell_images,1):
        files[f'xl/media/image{i}.png']=img
    struct="""<?xml version="1.0" encoding="UTF-8" standalone="yes"?><rvStructures xmlns="http://schemas.microsoft.com/office/spreadsheetml/2017/richdata" count="1"><s t="_localImage"><k n="_rvRel:LocalImageIdentifier" t="i"/><k n="CalcOrigin" t="i"/></s></rvStructures>"""
    rv=f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?><rvData xmlns="http://schemas.microsoft.com/office/spreadsheetml/2017/richdata" count="{len(cell_images)}">"""
    rv+=''.join(f'<rv s="0"><v>{i}</v><v>5</v></rv>' for i in range(len(cell_images)))+'</rvData>'
    rvr="""<?xml version="1.0" encoding="UTF-8" standalone="yes"?><richValueRels xmlns="http://schemas.microsoft.com/office/spreadsheetml/2022/richvaluerel" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">"""
    rvr+=''.join(f'<rel r:id="rId{i+1}"/>' for i in range(len(cell_images)))+'</richValueRels>'
    types="""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<rvTypesInfo xmlns="http://schemas.microsoft.com/office/spreadsheetml/2017/richdata2" xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" mc:Ignorable="x" xmlns:x="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><global><keyFlags><key name="_Self"><flag name="ExcludeFromFile" value="1"/><flag name="ExcludeFromCalcComparison" value="1"/></key><key name="_DisplayString"><flag name="ExcludeFromCalcComparison" value="1"/></key><key name="_Flags"><flag name="ExcludeFromCalcComparison" value="1"/></key><key name="_Format"><flag name="ExcludeFromCalcComparison" value="1"/></key><key name="_SubLabel"><flag name="ExcludeFromCalcComparison" value="1"/></key><key name="_Attribution"><flag name="ExcludeFromCalcComparison" value="1"/></key><key name="_Icon"><flag name="ExcludeFromCalcComparison" value="1"/></key><key name="_Display"><flag name="ExcludeFromCalcComparison" value="1"/></key><key name="_CanonicalPropertyNames"><flag name="ExcludeFromCalcComparison" value="1"/></key><key name="_ClassificationId"><flag name="ExcludeFromCalcComparison" value="1"/></key></keyFlags></global></rvTypesInfo>"""
    future_bks = ''.join(f'<bk><extLst><ext uri="{{3e2802c4-a4d2-4d8b-9148-e3be6c30e623}}"><xlrd:rvb i="{i}"/></ext></extLst></bk>' for i in range(len(cell_images)))
    value_bks = ''.join(f'<bk><rc t="1" v="{i}"/></bk>' for i in range(len(cell_images)))
    metadata=f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?><metadata xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:xlrd="http://schemas.microsoft.com/office/spreadsheetml/2017/richdata"><metadataTypes count="1"><metadataType name="XLRICHVALUE" minSupportedVersion="120000" copy="1" pasteAll="1" pasteValues="1" merge="1" splitFirst="1" rowColShift="1" clearFormats="1" clearComments="1" assign="1" coerce="1"/></metadataTypes><futureMetadata name="XLRICHVALUE" count="{len(cell_images)}">{future_bks}</futureMetadata><valueMetadata count="{len(cell_images)}">{value_bks}</valueMetadata></metadata>"""

    files['xl/richData/rdrichvaluestructure.xml']=struct.encode()
    files['xl/richData/rdrichvalue.xml']=rv.encode()
    files['xl/richData/richValueRel.xml']=rvr.encode()
    files['xl/richData/rdRichValueTypes.xml']=types.encode()
    files['xl/metadata.xml']=metadata.encode()
    relxml='<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'+''.join(f'<Relationship Id="rId{i+1}" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" Target="../media/image{i+1}.png"/>' for i in range(len(cell_images)))+'</Relationships>'
    files['xl/richData/_rels/richValueRel.xml.rels']=relxml.encode()
    wr=files['xl/_rels/workbook.xml.rels'].decode('utf-8')
    ids=[int(x) for x in re.findall(r'Id="rId(\d+)"',wr)]
    nxt=max(ids or [0])+1
    additions=(f'<Relationship Id="rId{nxt}" Type="http://schemas.microsoft.com/office/2017/06/relationships/rdRichValueStructure" Target="richData/rdrichvaluestructure.xml"/><Relationship Id="rId{nxt+1}" Type="http://schemas.microsoft.com/office/2017/06/relationships/rdRichValue" Target="richData/rdrichvalue.xml"/><Relationship Id="rId{nxt+2}" Type="http://schemas.microsoft.com/office/2022/10/relationships/richValueRel" Target="richData/richValueRel.xml"/><Relationship Id="rId{nxt+3}" Type="http://schemas.microsoft.com/office/2017/06/relationships/rdRichValueTypes" Target="richData/rdRichValueTypes.xml"/><Relationship Id="rId{nxt+4}" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/sheetMetadata" Target="metadata.xml"/>')
    files['xl/_rels/workbook.xml.rels']=wr.replace('</Relationships>',additions+'</Relationships>').encode()
    ct=files['[Content_Types].xml'].decode('utf-8')
    if '<Default Extension="png"' not in ct:
        ct=ct.replace('<Types ', '<Types ').replace('>', '><Default Extension="png" ContentType="image/png"/>', 1)
    additions=('<Override PartName="/xl/metadata.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheetMetadata+xml"/><Override PartName="/xl/richData/richValueRel.xml" ContentType="application/vnd.ms-excel.richvaluerel+xml"/><Override PartName="/xl/richData/rdrichvalue.xml" ContentType="application/vnd.ms-excel.rdrichvalue+xml"/><Override PartName="/xl/richData/rdrichvaluestructure.xml" ContentType="application/vnd.ms-excel.rdrichvaluestructure+xml"/><Override PartName="/xl/richData/rdRichValueTypes.xml" ContentType="application/vnd.ms-excel.richvaluetypes+xml"/>')
    files['[Content_Types].xml']=ct.replace('</Types>',additions+'</Types>').encode()
    outb=io.BytesIO()
    with zipfile.ZipFile(outb,'w',zipfile.ZIP_DEFLATED) as zout:
        for n,b in files.items(): zout.writestr(n,b)
    return outb.getvalue()


def export_xlsx(df, uploaded_files=None, exclude_legend=True):
    """XLSX export with screenshots only for blocks present in the legend, fitted to cells."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    import os

    out_df = df.copy() if isinstance(df, pd.DataFrame) else pd.DataFrame()
    out_df = apply_legend_names_to_results(out_df, st.session_state.get('legend_maps', {}))
    out_df = enrich_legend_names_from_files(out_df, uploaded_files=uploaded_files, exclude_legend=exclude_legend)
    for col in RESULT_COLUMNS:
        if col not in out_df.columns:
            out_df[col] = ''
    out_df = out_df[RESULT_COLUMNS]
    out_df['Nazwa z legendy'] = out_df['Nazwa z legendy'].replace({
        'gniazdo pojedyńcze': 'gniazdo pojedyncze',
        'dwa gniazda pojedyńcze RJ45, IP20': 'dwa gniazda pojedyncze RJ45, IP20',
    }, regex=True)
    out_df['Ilość'] = pd.to_numeric(out_df['Ilość'], errors='coerce').fillna(0).astype(int)

    summary = (out_df.groupby(['Typ elementu', 'Misc', 'Name', 'Nazwa elementu', 'Nazwa z legendy', 'Nazwa bloku'], as_index=False)['Ilość']
               .sum().sort_values(['Typ elementu', 'Nazwa elementu', 'Nazwa bloku'])) if not out_df.empty else pd.DataFrame(columns=['Typ elementu','Misc','Name','Nazwa elementu','Nazwa z legendy','Nazwa bloku','Ilość'])
    file_type = (out_df.groupby(['Plik', 'Typ elementu'], as_index=False)['Ilość']
                 .sum().sort_values(['Plik','Ilość'], ascending=[True,False])) if not out_df.empty else pd.DataFrame(columns=['Plik','Typ elementu','Ilość'])

    wb=Workbook(); wb.remove(wb.active)
    ws_sum=wb.create_sheet('Podsumowanie'); ws_pos=wb.create_sheet('Pozycje')
    ws_ft=wb.create_sheet('Plik_typ')

    def write_df(ws,data):
        ws.append(list(data.columns))
        for row in data.itertuples(index=False,name=None): ws.append(list(row))

    # Pozycje: eksportujemy tylko kolumny widoczne w docelowym zestawieniu.
    # Pola techniczne Misc/Name/Nazwa elementu pozostają wewnętrznie do mapowania
    # legendy i grafiki, ale nie są prezentowane w arkuszu Pozycje.
    pozycje_cols = [
        'Plik', 'Strona', 'Typ elementu', 'Nazwa bloku', 'Nazwa z legendy',
        'Symbol/etykieta', 'Ilość', 'Metoda', 'Pewność', 'Źródło'
    ]
    pozycje_export = out_df.copy()
    for col in pozycje_cols:
        if col not in pozycje_export.columns:
            pozycje_export[col] = ''
    pozycje_export = pozycje_export[pozycje_cols].copy()
    write_df(ws_pos,pozycje_export)
    # Podsumowanie: 6 kolumn; grafika wyłącznie w kolumnie F.
    summary_export=summary.copy()
    for col in ['Nazwa z legendy','Typ elementu','Nazwa elementu','Nazwa bloku','Ilość']:
        if col not in summary_export.columns:
            summary_export[col]=''
    summary_export=summary_export[
        ['Nazwa z legendy','Typ elementu','Nazwa elementu','Nazwa bloku','Ilość']
    ].copy()
    summary_export['Oznaczenie graficzne']=''
    summary_export.columns=[
        'Nazwa z legendy','Typ elementu','Nazwa elementu',
        'Nazwa bloku','Ilość','Oznaczenie graficzne'
    ]
    write_df(ws_sum,summary_export)
    write_df(ws_ft,file_type)

    header_fill=PatternFill('solid',fgColor='1F4E78'); header_font=Font(color='FFFFFF',bold=True); thin=Side(style='thin',color='D9E1F2')
    for ws in wb.worksheets:
        ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
        for cell in ws[1]:
            cell.fill=header_fill; cell.font=header_font
            cell.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
            cell.border=Border(bottom=thin)
        for col_cells in ws.columns:
            letter=get_column_letter(col_cells[0].column)
            max_len=max((len(str(c.value)) if c.value is not None else 0) for c in col_cells)
            ws.column_dimensions[letter].width=min(max(max_len+2,10),42)
        ws.row_dimensions[1].height=24

    for ws_name in ('Pozycje','Podsumowanie','Plik_typ'):
        ws=wb[ws_name]
        headers={str(c.value): c.column for c in ws[1] if c.value}
        qty_col=headers.get('Ilość')
        if qty_col:
            for row in range(2,ws.max_row+1): ws.cell(row,qty_col).number_format='0'

    wanted=set()
    if not summary.empty:
        for _, r in summary.iterrows():
            for key in (r.get('Nazwa elementu',''), r.get('Name',''), r.get('Misc',''), r.get('Nazwa bloku','')):
                key=str(key or '').strip()
                if key:
                    wanted.add(key)
    graphics={}
    try: graphics=collect_legend_graphics(uploaded_files or [], desired_blocks=wanted, exclude_legend=exclude_legend)
    except Exception: graphics={}

    temp_paths=[]
    cell_images=[]
    for ws_name, image_col, max_w, max_h, row_h in [
        ('Podsumowanie','F',105,62,68),
    ]:
        ws=wb[ws_name]
        ws.column_dimensions[image_col].width=18
        # Image cell is centered visually by using alignment on the anchor cell.
        for ridx in range(2,ws.max_row+1):
            element_name=''
            block_name=''
            headers={str(c.value): c.column for c in ws[1] if c.value}
            if 'Nazwa elementu' in headers:
                element_name=str(ws.cell(ridx,headers['Nazwa elementu']).value or '')
            if 'Nazwa bloku' in headers:
                block_name=str(ws.cell(ridx,headers['Nazwa bloku']).value or '')
            name_attr=str(ws.cell(ridx,headers['Name']).value or '') if 'Name' in headers else ''
            misc_attr=str(ws.cell(ridx,headers['Misc']).value or '') if 'Misc' in headers else ''
            block=(name_attr.strip() or misc_attr.strip() or element_name.strip() or block_name.strip())
            ws.cell(ridx,ws[image_col+'1'].column).alignment=Alignment(horizontal='center',vertical='center')
            ws.cell(ridx,ws[image_col+'1'].column).fill=PatternFill('solid',fgColor='FFFFFF')
            img_bytes=graphics.get(block)
            if not img_bytes:
                continue
            cell_images.append((f'{image_col}{ridx}', img_bytes))
        # Make the image column explicitly the graphic column.
        ws.cell(1,ws[image_col+'1'].column).value='Oznaczenie graficzne'
        ws.cell(1,ws[image_col+'1'].column).fill=header_fill; ws.cell(1,ws[image_col+'1'].column).font=header_font

    out=io.BytesIO()
    try:
        wb.save(out)
    finally:
        for p in temp_paths:
            try: os.unlink(p)
            except Exception: pass
    return _convert_graphics_to_incell_xlsx(out.getvalue(), cell_images)


def project_json(files_meta):
    return json.dumps({
        'version': APP_VERSION,
        'project': st.session_state.project_name,
        'files': files_meta,
        'results': st.session_state.results.to_dict(orient='records'),
        'manual': st.session_state.manual,
        'patterns': st.session_state.patterns,
        'dxf_count_mode': st.session_state.dxf_count_mode,
    }, ensure_ascii=False, indent=2).encode('utf-8')



def build_route_measurement_exports(measurements, summary_columns=None):
    """Create deterministic XLSX/detail CSV/summary CSV payloads for route measurements.

    Returns a plain dict of bytes suitable for Streamlit download buttons. The workbook
    always contains two visible worksheets, even when one of the data frames is empty.
    """
    import io
    import pandas as pd
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter

    md = measurements.copy() if isinstance(measurements, pd.DataFrame) else pd.DataFrame()
    if summary_columns is None:
        summary_columns = ['Koryto', 'Szerokość [mm]', 'Długość [m]', 'Obiekty', 'Jednostka', 'Ilość [szt.]']
    missing = [c for c in summary_columns if c not in md.columns and c not in {'Jednostka', 'Ilość [szt.]'}]
    if missing:
        raise ValueError(f"Brak wymaganych kolumn pomiaru: {', '.join(missing)}")

    # Keep countable components (e.g. KK_S_* holders) separate from linear metres.
    if 'Jednostka' in md.columns:
        md['Jednostka'] = md['Jednostka'].fillna('m')
    else:
        md['Jednostka'] = 'm'
    if 'Ilość [szt.]' in md.columns:
        md['Ilość [szt.]'] = pd.to_numeric(md['Ilość [szt.]'], errors='coerce').fillna(0)
    else:
        md['Ilość [szt.]'] = 0
    group_cols = ['Koryto', 'Szerokość [mm]', 'Jednostka']
    summary_all = md.groupby(group_cols, dropna=False, as_index=False).agg(**{
        'Długość [m]': ('Długość [m]', 'sum'),
        'Obiekty': ('Obiekty', 'sum'),
        'Ilość [szt.]': ('Ilość [szt.]', 'sum'),
    })
    for col in summary_columns:
        if col not in summary_all.columns:
            summary_all[col] = ''
    summary = summary_all[summary_columns]

    wb = Workbook()
    ws_summary = wb.active
    ws_summary.title = 'Koryta wg szerokości'
    ws_detail = wb.create_sheet('Pomiar tras')
    for ws, frame in ((ws_summary, summary), (ws_detail, md)):
        ws.sheet_state = 'visible'
        ws.freeze_panes = 'A2'
        ws.append([str(c) for c in frame.columns])
        for row in frame.itertuples(index=False, name=None):
            ws.append([None if pd.isna(v) else v.item() if hasattr(v, 'item') else v for v in row])
        for idx, column in enumerate(ws.iter_cols(), start=1):
            max_len = max((len(str(cell.value or '')) for cell in column), default=0)
            ws.column_dimensions[get_column_letter(idx)].width = max(12, min(35, max_len + 2))
    for ws in wb.worksheets:
        ws.sheet_state = 'visible'
    wb.active = 0
    xlsx = io.BytesIO()
    wb.save(xlsx)
    return {
        'xlsx': xlsx.getvalue(),
        'detail_csv': md.to_csv(index=False, sep=';', encoding='utf-8-sig', lineterminator='\n').encode('utf-8-sig'),
        'summary_csv': summary.to_csv(index=False, sep=';', encoding='utf-8-sig', lineterminator='\n').encode('utf-8-sig'),
    }
