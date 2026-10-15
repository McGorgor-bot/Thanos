import io
from openpyxl import load_workbook
import pandas as pd
from thanos.cable_list import _circuit_phase_counts, export_cable_list_xlsx, CABLE_COLUMNS


def test_phase_classification_core_counts():
    assert _circuit_phase_counts('N2XH-J 3x1,5 mm²') == (1, 0)
    assert _circuit_phase_counts('N2XH-J 5x6 mm²') == (0, 1)
    assert _circuit_phase_counts('4x N2HX 1x120 + N2XH 1x70') == (0, 1)
    assert _circuit_phase_counts('YDY 2x1,5') == (1, 0)
    assert _circuit_phase_counts('unknown cable') == (0, 0)


def test_xlsx_contains_phase_columns_and_summary_counts():
    rows = [
        {'Typ kabla': 'N2XH-J 3x1,5 mm²', 'Liczba równoległych kabli': 1, 'Długość [m]': None},
        {'Typ kabla': 'N2XH-J 5x6 mm²', 'Liczba równoległych kabli': 1, 'Długość [m]': None},
        {'Typ kabla': '4x N2HX 1x120 + N2XH 1x70', 'Liczba równoległych kabli': 4, 'Długość [m]': None},
    ]
    df = pd.DataFrame(rows, columns=CABLE_COLUMNS)
    wb = load_workbook(io.BytesIO(export_cable_list_xlsx(df)), data_only=True)
    ws = wb['Lista kablowa']
    headers = [c.value for c in ws[1]]
    assert headers[-2:] == ['obwód 1-f', 'obwód 3-f']
    idx1, idx3 = headers.index('obwód 1-f') + 1, headers.index('obwód 3-f') + 1
    assert [(ws.cell(r, idx1).value, ws.cell(r, idx3).value) for r in range(2, 5)] == [(1, 0), (0, 1), (0, 1)]
    summary = wb['Podsumowanie typów']
    sh = [c.value for c in summary[1]]
    assert sh[-2:] == ['obwód 1-f', 'obwód 3-f']
